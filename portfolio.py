from systemMarket import systemMarketClass
from turtleEquity import plotEquityCurve
from math import sqrt
try:
    from openpyxl import Workbook
    from openpyxl.chart.axis import DateAxis
    from openpyxl.styles import NamedStyle
    from openpyxl.chart import (
        LineChart,
        Reference,
        Series,
        )
    exportToXL = True
except ImportError:
    exportToXL = False

from datetime import date

class portfolioClass(object):
    def __init__(self):
        self.portfolioName = ""
        self.systemMarkets = list()
        self.portEquityDate = list()
        self.portEquityVal = list()
        self.portclsTrdEquity = list()
        self.portDailyEquityVal = list()
        self.portPeakEquity = 0
        self.portMinEquity = 0
        self.portMaxDDPer = 0.0
        self.portMaxDD = 0
        self.avgMonthlyReturn=0
        self.monthlyStdDev = 0
        self.avgYearlyReturn = 0
        self.avgMaxDDTuple = list()
        self.avgMaxDD = 0
        self.portTrades= 0
        self.yearlyStdDev = 0
        tempEqu = 0
        cumEqu = 0
        maxEqu = -999999999
        minEqu = 999999999
        maxDD = 0

    def setPortfolioInfo(self,name,systemMarket):
        self.portfolioName = name
        self.systemMarkets = list(systemMarket)
        masterDateList = list()
        monthList = list()
        monthEquity = list()
        yearEquity= list()
        combinedEquity = list()
        self.portPeakEquity = -999999999999
        self.portMinEquity = -999999999999
        monthToday = 0
        monthYesDay = 0
        newYear = True

        printOutDailyEquity = False
        printMonthlyToTerminal = False
#        exportToXL = True

        tempList = list()
        tempList.append(['Date','Cumulative'])


        fileName1 = self.systemMarkets[0].systemName + "-Composite.txt"
        target1 = open(fileName1,"w")
        fileName2 = self.systemMarkets[0].systemName + "-DailyEquity.txt"
        if printOutDailyEquity :
            target2 = open(fileName2,"w")
        fileName3 = self.systemMarkets[0].systemName + "-DrawDown.txt"
        target3 = open(fileName3,"w")
        if exportToXL:
            xlFileName =self.systemMarkets[0].systemName+".xlsx"
            wb = Workbook()
            sheet = wb.create_sheet("EquityCurve")
            sheet = wb.get_sheet_by_name("EquityCurve")
            date_style = NamedStyle(name = 'datetime',number_format='MM/DD/YYYY')
            sheet['A1'].style = date_style
#            sheet = wb.active
#            sheet = wb.
#            for row in tempList:
#                    sheet.append(row)
#            wb.save(xlFileName)

        begDate = 99999999
        endDate = 0
        for i in range(0,len(self.systemMarkets)):
            if(self.systemMarkets[i].equity.equityDate[0] < begDate):
                begDate = self.systemMarkets[i].equity.equityDate[0]
            if(self.systemMarkets[i].equity.equityDate[-1] > endDate):
               endDate = self.systemMarkets[i].equity.equityDate[-1]

        for i in range(0,len(self.systemMarkets)):
            masterDateList += self.systemMarkets[i].equity.equityDate
            sysName = self.systemMarkets[i].systemName
            market = self.systemMarkets[i].symbol
            comName = self.systemMarkets[i].comName
            avgWin = self.systemMarkets[i].avgWin
            sysMark =self.systemMarkets[i]
            avgLoss = sysMark.avgLoss
            totProf = sysMark.profitLoss
            totTrades = sysMark.numTrades
            self.portTrades += totTrades
            avgMonthlyRet = sysMark.avgMonthlyReturn
            monthlyStdDev = sysMark.avgMonthlyStdDev
            avgYearlyRet = sysMark.avgYearlyReturn
            maxD = sysMark.maxxDD
            clsTrdDD = sysMark.clsTrdDD
            perWins = sysMark.perWins
            tempStr =''
            if len(sysName) < 15 :
                for j in range(0,15 - len(sysName)):
                    sysName = sysName + ' '
            if len(sysName) > 15: sysName = sysName[0:12]+'...'
            if i == 0:
                print('                                                                                      Avg.    Monthly    Avg.')
                print('SysName         Market     TotProfit  MaxDD ClsTrdDD AvgWin  AvgLoss PerWins #Trds  MonthRet.  StdDev  YearlyRet.')
                print('-------------------------------------------------------------------------------------------------------------------')
                lineOutPut = 'Testing from : ' + str(begDate) + ' to: ' + str(endDate) + '\n'
                target1.write(lineOutPut)
                lineOutPut = '-----------------------------------------------------------------------------------------------------------\n'
                target1.write(lineOutPut)
                lineOutPut = '                                                                                      Avg.   Monthly    Avg.\n'
                target1.write(lineOutPut)
                lineOutPut = 'SysName         Market     TotProfit  MaxDD ClsTrdDD AvgWin  AvgLoss PerWins #Trds  MonthRet.  StdDev  YearlyRet.\n'
                target1.write(lineOutPut)
            print('%s %-10s %9d %6d  %6d  %6d   %6d    %3.2f   %4d    %5d   %5d   %6d' % (sysName,comName+"-"+market,totProf,maxD,clsTrdDD,avgWin,avgLoss,perWins,totTrades,avgMonthlyRet,monthlyStdDev,avgYearlyRet))
            target1.write('%s %-10s %9d %6d  %6d  %6d   %6d    %3.2f   %4d    %5d   %5d   %6d\n' % (sysName,comName+"-"+market,totProf,maxD,clsTrdDD,avgWin,avgLoss,perWins,totTrades,avgMonthlyRet,monthlyStdDev,avgYearlyRet))
        print('-------------------------------------------------------------------------------------------------------------------')
        lineOutPut = '-------------------------------------------------------------------------------------------------------------------\n'
        target1.write(lineOutPut)
        masterDateList = removeDuplicates(masterDateList)
        masterDateList = sorted(masterDateList)
        self.portEquityDate = masterDateList
        monthList = createMonthList(masterDateList)
        pEquityTuple = list()
        self.drawDownTuple = list()
        numDaysInMList = len(masterDateList)
        boyDrawDown = 0
        ddCnt = 0
        temp11 = 0
        dailyPortEquity = 0
        priorIdxlist = list()
        indivMktEquList = list()
        numMarkets = len(self.systemMarkets)
        priorIdxList = [0] * len(self.systemMarkets)
        for i in range(0,len(masterDateList)):
            cumuVal = 0
            portDailyEquity = 0
            if i > 0:
                monthYesDay = (int(masterDateList[i-1]/100) % 100)
                monthToDay =(int(masterDateList[i]/100) % 100)
            idx = 0
            indivMktEquList = [0] * len(self.systemMarkets)
            for j in range(0,len(self.systemMarkets)):
                skipDay = 0
                try:
                    idx = self.systemMarkets[j].equity.equityDate.index(masterDateList[i])
                except ValueError:
                    skipDay = 1
                    marketDayCumu = 0
                    skipDate = masterDateList[i]
                    skipMkt = self.systemMarkets[j].symbol
                    numDaysInEquityStream = len(self.systemMarkets[j].equity.dailyEquityVal)
                    marketBeginDate = self.systemMarkets[j].equity.equityDate[0]
                    if (masterDateList[i] > marketBeginDate and i < numDaysInEquityStream):
                        marketDayCumu = self.systemMarkets[j].equity.dailyEquityVal[priorIdxList[j]]
                    else:
                        if masterDateList[i] < marketBeginDate:
                            marketDayCumu = 0
                        else:
                            marketDayCumu = self.systemMarkets[j].equity.dailyEquityVal[-1]
                    indivMktEquList[j] = marketDayCumu
                    pEquityTuple += ((i,j,marketDayCumu),)
                    portDailyEquity += marketDayCumu
                    cumuVal += marketDayCumu
#                    print("Day Skipped ",self.systemMarkets[j].symbol," ",masterDateList[i]," ",marketDayCumu," ",cumuVal," ",priorIdxList[j])
                if skipDay == 0:
                    priorIdxList[j] = idx
                    marketDayCumu = self.systemMarkets[j].equity.dailyEquityVal[idx]
                    pEquityTuple += ((i,j,marketDayCumu),)
                    indivMktEquList[j] = marketDayCumu
                    portDailyEquity += marketDayCumu
                    cumuVal += self.systemMarkets[j].equity.dailyEquityVal[idx]
##                    if masterDateList[i] == 20080118 :
##                        print("Day 20080118 ",self.systemMarkets[j].symbol," ",masterDateList[i]," ",marketDayCumu," ",cumuVal," ",idx)
##
##                    if masterDateList[i] == 20080121 :
##                        print("Day 20080121 ",self.systemMarkets[j].symbol," ",masterDateList[i]," ",marketDayCumu," ",cumuVal)
            if printOutDailyEquity :
                target2.write('%8d %8.2f ' % (masterDateList[i],portDailyEquity))
                for x in range(0,numMarkets):
                    target2.write('%8.2f ' % (indivMktEquList[x]))
                target2.write('\n')
            combinedEquity.append(cumuVal)
            self.portEquityVal.append(cumuVal)
            if newYear :
                temp11 = cumuVal
            newYear = False
            temp11 = max(temp11,cumuVal)
            boyDrawDown = max(boyDrawDown,temp11 - cumuVal)
            if (i > 0 and monthYesDay == 12 and monthToDay == 1) or i == len(masterDateList) - 1:
                self.drawDownTuple += ((ddCnt,boyDrawDown),)
#                print(masterDateList[i]," ",boyDrawDown)
                boyDrawDown = 0
                newYear = True
                ddCnt += 1
            if cumuVal > self.portPeakEquity:
                self.portPeakEquity = cumuVal
                drawDownChart = 0
            self.portMinEquity = max(self.portMinEquity,self.portPeakEquity - cumuVal)
            target3.write('%8d %8.2f\n' % (masterDateList[i],cumuVal - self.portPeakEquity))
            self.portMaxDD = self.portMinEquity
            if self.portPeakEquity != 0 : self.portMaxDDPer = self.portMaxDD/self.portPeakEquity*100

        numYearlyDD = len(self.drawDownTuple)
        tempYearlyDDList = ([x[1] for x in self.drawDownTuple])
        mean = sum(tempYearlyDDList)/len(tempYearlyDDList)
        self.avgMaxDD = mean
        tempStr = tempStr + '   '
        print('Totals                  %12d %6d  %4.1f%% Avg. DD %7d %6d' % (self.portEquityVal[-1],self.portMaxDD,self.portMaxDDPer,self.avgMaxDD,self.portTrades))
        lineOutPut = 'Totals                  '
        lineOutPut2 ='% Avg. DD'
        lineOutPut3 = "    "
        lineOutPut4 = "or"
        target1.write('%-s %12d %6d %-s %4.1f %-s %7d %-s %6d\n' % (lineOutPut,self.portEquityVal[-1],self.portMaxDD,lineOutPut4,self.portMaxDDPer,lineOutPut2,self.avgMaxDD,lineOutPut3,self.portTrades))
        print('-------------------------------------------------------------------')
        lineOutPut = '-----------------------------------------------------------------------------------------------------------\n'
        target1.write(lineOutPut)

##        print("Combined Equity: ",self.portEquityVal[-1])
##        lineOutPut = "Combined Equity: "
##        target1.write('%s %7.3f\n' %(lineOutPut,self.portEquityVal[-1]))
##        print("Combined MaxDD: ",self.portMaxDD)
##        lineOutPut = "Combined MaxDD: "
##        target1.write('%s %7.3f\n' %(lineOutPut,self.portMaxDD))
        if printMonthlyToTerminal: print("Combined Monthly Return")
        if printMonthlyToTerminal: print("Date         Profit Cum.Profit")
        lineOutPut = "Combined Monthly Return\n"
        target1.write(lineOutPut)
        lineOutPut = "Date         Profit Cum.Profit\n"
        target1.write(lineOutPut)
        if printMonthlyToTerminal: print('-------------------------------------------------------------------')
        lineOutPut = '-------------------------------------------------------------------\n'
        target1.write(lineOutPut)

        begOfYearEquity = 0
        numOfMonths = len(monthList)
        for j in range(0,len(monthList)):
            idx = masterDateList.index(monthList[j])
            combinedDailyEquity = 0
            if j == 0:
                monthEquity.append(combinedEquity[idx])
                prevCombinedDailyEquity = monthEquity[-1]
            else:
                combinedDailyEquity = combinedEquity[idx]
                monthEquity.append(combinedDailyEquity -  prevCombinedDailyEquity)
                prevCombinedDailyEquity = combinedDailyEquity
            isDecember = False
            if(int(monthList[j]/100) % 100) == 12 or (j == numOfMonths-1):
                isDecember = True
                yearEquity.append(combinedDailyEquity - begOfYearEquity)
                begOfYearEquity = combinedDailyEquity
            if exportToXL:
                yr =  int(monthList[j]/10000)
                mn =  int((monthList[j]/100)%100)
                dy = monthList[j]%100
                tempList.append([date(yr,mn,dy),combinedEquity[idx]])


            if isDecember == False:
                if printMonthlyToTerminal: print('%8d %10.0f %10.0f ' % (monthList[j],monthEquity[j],combinedEquity[idx]))
                target1.write('%8d %10.0f %10.0f\n' % (monthList[j],monthEquity[j],combinedEquity[idx]))
            if isDecember == True:
                if printMonthlyToTerminal: print('%8d %10.0f %10.0f %10.0f' % (monthList[j],monthEquity[j],combinedEquity[idx],yearEquity[-1]))
                target1.write('%8d %10.0f %10.0f %10.0f\n' % (monthList[j],monthEquity[j],combinedEquity[idx],yearEquity[-1]))
        if exportToXL:
            for rows in tempList:
                sheet.append(rows)
#            refObj = Reference(sheet,min_col=1,min_row = 2,max_col=2,max_row=len(tempList))
#            seriesObj = Series(refObj, title='Cumu Equity')
##            chartObj = LineChart()
##            chartObj.title = 'Equity Chart'
###            chartObj.style = 13
###            chartObj.append(seriesObj)
##            chartObj.x_axis.title = 'Date'
##            chartObj.y_axis.title = 'Equity'
###            chartObj.x_axis = DateAxis(crossAx = 10000)
##            chartObj.x_axis.number_format = 'yyyy/mm/dd'
##            chartObj.x_axis.majorTimeUnit = 'days'
##            data = Reference(sheet,min_col=2,min_row=3,max_row=len(tempList))
##            chartObj.add_data(data,titles_from_data = False)
##            dates = Reference(sheet,min_col=1,min_row=3,max_row=len(tempList))
##            chartObj.set_categories(dates)
            c1 = LineChart()
            c1.x_axis = DateAxis() # axId defaults to 500
            c1.x_axis.title = "Date"
            c1.x_axis.crosses = "min"
            c1.x_axis.majorTickMark = "out"
            c1.x_axis.number_format = "yy-mm-dd"
            data = Reference(sheet,min_col=2,min_row=2,max_row=len(tempList))
            dates = Reference(sheet,min_col=1,min_row=3,max_row=len(tempList))
            c1.add_data(data, titles_from_data=False)
            c1.set_categories(dates)
            c1.y_axis.title = 'Cumu. Equity'
            c1.y_axis.crossAx = 500
            c1.y_axis.majorGridlines = None
            sheet.add_chart(c1,'C5')
            wb.save(xlFileName)

        mean = sum(monthEquity) / len(monthEquity)
        differences = [x - mean for x in monthEquity]
        sq_differences = [d ** 2 for d in differences]
        ssd = sum(sq_differences)
        variance = ssd / ( len(monthEquity) - 1)
        sd = sqrt(variance)
        self.avgMonthlyReturn = mean
        self.monthStdDev = sd
        if printMonthlyToTerminal: print('Average Monthly Return: %8d StdDev: %5d  Sharpe: %4.2f' % (mean,sd,mean/sd))
        target1.write('Average Monthly Return: %8d StdDev: %5d  Sharpe: %4.2f\n' % (mean,sd,mean/sd))
        mean = sum(yearEquity) / len(yearEquity)
        differences = [x - mean for x in yearEquity]
        sq_differences = [d ** 2 for d in differences]
        ssd = sum(sq_differences)
        variance = ssd / ( len(yearEquity) - 1)
        sd = sqrt(variance)
        self.avgYearlyReturn = mean
        self.yearlyStdDev = sd
        if printMonthlyToTerminal: print('Average  Yearly Return: %8d StdDev: %5d  Sharpe: %4.2f' % (mean,sd,mean/sd))
        target1.write('Average  Yearly Return: %8d StdDev: %5d  Sharpe: %4.2f\n' % (mean,sd,mean/sd))
        target1.close()
        if printOutDailyEquity: target2.close()
        target3.close()
##        plotEquityCurve(self)
##        return(2)

def removeDuplicates(li):
    my_set = set()
    res = []
    for e in li:
        if e not in my_set:
            res.append(e)
            my_set.add(e)
    return res

def createMonthList(li):
    myMonthList = list()
    for i in range(0,len(li)):
        if i != 0:
            tempa = int(li[i]/100)
            pMonth = int(li[i-1]/100) % 100
            month = int(li[i]/100) % 100
            if pMonth != month:
                myMonthList.append(li[i-1])
            if i == len(li)-1:
                myMonthList.append(li[i])
    return myMonthList